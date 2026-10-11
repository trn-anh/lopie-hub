"""Tạo file lý lịch khoa học mẫu ĐÃ ĐIỀN (dữ liệu minh họa) để kiểm thử bộ tách CV.

Bố cục bám mẫu Bộ GD&ĐT (TT 08/2011, Phụ lục V): nhãn và giá trị nằm chung dòng,
ngăn bằng tab; các mục III, IV là bảng có hàng tiêu đề.

Chạy: python3 qtda/scholarlink/tools/make-fixture.py
Kết quả: qtda/scholarlink/tools/fixtures/llkh-mau-da-dien.docx
"""
import os
from docx import Document
from docx.shared import Pt

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "fixtures", "llkh-mau-da-dien.docx")

doc = Document()
st = doc.styles["Normal"]
st.font.name = "Times New Roman"
st.font.size = Pt(13)

P = doc.add_paragraph
P("CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM")
P("Độc lập – Tự do – Hạnh phúc")
P("LÝ LỊCH KHOA HỌC")
P("(Theo mẫu tại Thông tư số 08/2011/TT-BGDĐT ngày 17/02/2011 của Bộ trưởng Bộ GDĐT – Phụ lục V)")
P("I. LÝ LỊCH SƠ LƯỢC")
P("Họ và tên: Trần Minh Anh\tGiới tính: Nam")
P("Ngày, tháng, năm sinh: 15/08/1988\tNơi sinh: TP. Hồ Chí Minh")
P("Quê quán: Đồng Tháp\tDân tộc: Kinh")
P("Học vị cao nhất: Tiến sĩ\tNăm, nước nhận học vị: 2021, Việt Nam")
P("Chức danh khoa học cao nhất: ........................\tNăm bổ nhiệm: ..............")
P("Chức vụ (hiện tại hoặc trước khi nghỉ hưu): Phó Trưởng bộ môn")
P("Đơn vị công tác (hiện tại hoặc trước khi nghỉ hưu): Khoa Kinh tế, Trường ĐH Công nghệ Kỹ thuật TP.HCM")
P("Chỗ ở riêng hoặc địa chỉ liên lạc: 1 Võ Văn Ngân, phường Thủ Đức, TP. Hồ Chí Minh")
P("Điện thoại liên hệ:  CQ: 028 3896 8641\tNR: ........................\tDĐ: 0909 123 456")
P("Fax: ..............\tE-mail: anhtm@hcmute.edu.vn")
P("Số CMND: 079088001234\tNgày cấp: 12/03/2021\tNơi cấp: Cục Cảnh sát QLHC về TTXH")
P("II. QUÁ TRÌNH ĐÀO TẠO")
P("1. Đại học:")
P("Hệ đào tạo: Chính quy")
P("Nơi đào tạo: Trường ĐH Sư phạm Kỹ thuật TP.HCM")
P("Ngành học: Quản lý công nghiệp")
P("Nước đào tạo: Việt Nam\tNăm tốt nghiệp: 2010")
P("Bằng đại học 2: ..............................\tNăm tốt nghiệp: ............")
P("2. Sau đại học")
P("- Thạc sĩ chuyên ngành: Quản trị kinh doanh\tNăm cấp bằng: 2014")
P("Nơi đào tạo: Trường ĐH Kinh tế TP.HCM")
P("- Tiến sĩ chuyên ngành: Quản lý công nghiệp\tNăm cấp bằng: 2021")
P("Nơi đào tạo: Trường ĐH Bách khoa – ĐHQG TP.HCM")
P("- Tên luận án: Tối ưu hóa chuỗi cung ứng trong doanh nghiệp sản xuất vừa và nhỏ")

t = doc.add_table(rows=1, cols=3)
t.rows[0].cells[0].text = "3. Ngoại ngữ:"
t.rows[0].cells[1].text = "1. Tiếng Anh\n2. Tiếng Nhật\n3. …………………………"
t.rows[0].cells[2].text = "Mức độ sử dụng: IELTS 7.0\nMức độ sử dụng: N3\nMức độ sử dụng: ………………………"

P("III. QUÁ TRÌNH CÔNG TÁC CHUYÊN MÔN")
work = [("Thời gian", "Nơi công tác", "Công việc đảm nhiệm"),
        ("2010 – 2014", "Công ty TNHH Sản xuất An Phát", "Kỹ sư kế hoạch sản xuất"),
        ("2014 – nay", "Trường ĐH Công nghệ Kỹ thuật TP.HCM", "Giảng viên"),
        ("", "", "")]
t = doc.add_table(rows=0, cols=3)
for r in work:
    cells = t.add_row().cells
    for i, v in enumerate(r):
        cells[i].text = v

P("IV. QUÁ TRÌNH NGHIÊN CỨU KHOA HỌC")
P("1. Các đề tài nghiên cứu khoa học đã và đang tham gia:")
proj = [("TT", "Tên đề tài nghiên cứu", "Năm bắt đầu/Năm hoàn thành", "Đề tài cấp (NN, Bộ, ngành, trường)", "Trách nhiệm tham gia trong đề tài"),
        ("1", "Mô hình dự báo nhu cầu cho doanh nghiệp vừa và nhỏ", "2019/2021", "Trường", "Chủ nhiệm"),
        ("2", "Chuyển đổi số chuỗi cung ứng nông sản", "2022/2024", "Bộ", "Thành viên"),
        ("", "", "", "", "")]
t = doc.add_table(rows=0, cols=5)
for r in proj:
    cells = t.add_row().cells
    for i, v in enumerate(r):
        cells[i].text = v

P("2. Các công trình khoa học đã công bố:")
pubs = [("TT", "Tên công trình", "Năm công bố", "Tên tạp chí"),
        ("1", "Tối ưu chuỗi cung ứng trong sản xuất", "2026", "Tạp chí Khoa học Giáo dục Kỹ thuật"),
        ("2", "Chuyển đổi số trong doanh nghiệp vừa và nhỏ", "2026", "Kỷ yếu Hội nghị Quản lý Công nghiệp"),
        ("3", "Nguyen, H., & Tran, M. A. (2025). Multi-criteria decision model for supplier selection. Journal of Industrial Engineering, 12(3). https://doi.org/10.1234/jie.2025.0042", "2025", "Journal of Industrial Engineering"),
        ("", "", "", "")]
t = doc.add_table(rows=0, cols=4)
for r in pubs:
    cells = t.add_row().cells
    for i, v in enumerate(r):
        cells[i].text = v

P("TP. Hồ Chí Minh, ngày 10 tháng 10 năm 2026")
t = doc.add_table(rows=1, cols=2)
t.rows[0].cells[0].text = "Xác nhận của cơ quan"
t.rows[0].cells[1].text = "Người khai ký tên\n(Ghi rõ chức danh, học vị)"

os.makedirs(os.path.dirname(OUT), exist_ok=True)
doc.save(OUT)
print("Đã tạo", OUT)
